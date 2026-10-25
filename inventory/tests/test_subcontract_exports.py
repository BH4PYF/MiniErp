import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from ..models import (
    Budget, BudgetItem, Contract, ContractItem, Measurement, Profile,
    Project, SubcontractCategory, SubcontractList, Subcontractor,
)


class SubcontractExportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username='export_admin', password='pass12345'
        )
        Profile.objects.create(user=cls.user, role='admin')
        cls.project = Project.objects.create(code='PRJ01', name='项目A')
        cls.subcontractor = Subcontractor.objects.create(
            code='SUB01', name='分包商A', contact='李四',
            phone='13900000000', main_type='土建',
        )
        cls.category = SubcontractCategory.objects.create(
            category_code='SCAT01', category_name='土建',
        )
        cls.sublist = SubcontractList.objects.create(
            code='SL01', name='混凝土', category=cls.category,
            construction_params='C30', unit='m3',
        )
        cls.budget = Budget.objects.create(code='B-0001', project=cls.project)
        BudgetItem.objects.create(
            budget=cls.budget, item_order=1, subcontract_list=cls.sublist,
            quantity=Decimal('10'), unit_price=Decimal('100'),
        )
        cls.contract = Contract.objects.create(
            code='C-0001', name='合同A',
            project=cls.project, subcontractor=cls.subcontractor,
        )
        ContractItem.objects.create(
            contract=cls.contract, item_order=1, subcontract_list=cls.sublist,
            quantity=Decimal('20'), unit_price=Decimal('100'),
        )
        Measurement.objects.create(
            code='M-0001', contract=cls.contract, project=cls.project,
            subcontractor=cls.subcontractor,
            period_start=date(2026, 1, 1), period_end=date(2026, 1, 31),
            current_value=Decimal('500'),
        )

    def setUp(self):
        self.client.force_login(self.user)

    def load_sheet(self, response):
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        return workbook.active

    def test_contract_list_export_computes_totals(self):
        response = self.client.get(reverse('export_contracts'))
        sheet = self.load_sheet(response)

        self.assertEqual(sheet.cell(row=2, column=5).value, 2000.0)
        self.assertEqual(sheet.cell(row=2, column=6).value, 500.0)
        self.assertEqual(sheet.cell(row=2, column=7).value, 25.0)

    def test_budget_list_export_computes_totals(self):
        response = self.client.get(reverse('export_budgets'))
        sheet = self.load_sheet(response)

        self.assertEqual(sheet.cell(row=2, column=3).value, 1000.0)
        self.assertEqual(sheet.cell(row=2, column=4).value, 500.0)
        self.assertEqual(sheet.cell(row=2, column=5).value, 50.0)

    def test_contract_detail_export_writes_category_name(self):
        response = self.client.get(
            reverse('export_contract_detail', args=[self.contract.pk])
        )
        sheet = self.load_sheet(response)

        self.assertEqual(sheet.cell(row=9, column=3).value, '土建')

    def test_budget_detail_export_writes_category_name(self):
        response = self.client.get(
            reverse('export_budget_detail', args=[self.budget.pk])
        )
        sheet = self.load_sheet(response)

        self.assertEqual(sheet.cell(row=8, column=3).value, '土建')
